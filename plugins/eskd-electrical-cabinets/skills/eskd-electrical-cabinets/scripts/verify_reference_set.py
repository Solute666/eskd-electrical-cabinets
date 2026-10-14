#!/usr/bin/env python3
"""Check the frozen X2002 set against its actual PDFs/XLSX and source models.

This is a project-specific consistency check, not an electrical safety approval.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re

import fitz
from openpyxl import load_workbook


def norm(value):
    return re.sub(r'\s+', '', str('' if value is None else value)).replace('−', '-').replace('–', '-').replace('—', '-')


def rows(path, sheet):
    book = load_workbook(path, read_only=True, data_only=True)
    result = list(book[sheet].iter_rows(values_only=True))
    book.close()
    return result


def check(root):
    manifest = json.loads((root / 'set-manifest.json').read_text())
    docs = manifest['documents']
    model = json.loads((root / 'model/mount_model_R29.json').read_text())
    elec = json.loads((root / 'model/electrical_connections_R27.json').read_text())
    layout = json.loads((root / 'model/layout_positions_R31.json').read_text())
    errors = []
    pdf_text = {}
    evidence = {}
    for role, item in docs.items():
        file = root / item['path']
        if not file.is_file():
            errors.append(f'Missing document: {role}')
            continue
        if hashlib.sha256(file.read_bytes()).hexdigest() != item['sha256']:
            errors.append(f'Changed frozen document: {role}')
        if file.suffix == '.pdf':
            with fitz.open(file) as pdf:
                if len(pdf) != item['pages']:
                    errors.append(f'Wrong page count: {role}')
                pdf_text[role] = [p.get_text() for p in pdf]
                if any(not t.strip() for t in pdf_text[role]):
                    errors.append(f'Empty PDF page: {role}')

    def path(role):
        return root / docs[role]['path']

    spec = {r[0]: r for r in rows(path('specification'), 'Спецификация') if r and type(r[0]) is int}
    elements = {r[0]: r for r in rows(path('mounting_workbook'), 'Перечень') if r and type(r[0]) is int}
    supply = {r[0]: r for r in rows(path('equipment_workbook'), 'Поставка') if r and type(r[0]) is int}
    if len(spec) != len(elements) or len(spec) != len(supply) or len(spec) != len(model['elements']):
        errors.append('Element/specification/supply row counts differ')
    for e in model['elements']:
        n = e['spec']
        wanted = [e['name'], e['model'], e['brand'], e['purchase'], e['unit']]
        actual = list(spec.get(n, [None] * 6)[1:6])
        if [norm(v) for v in actual] != [norm(v) for v in wanted]:
            errors.append(f'Specification row {n} differs from model')
        erow = elements.get(n, [None] * 12)
        actual = [erow[2], erow[3], erow[4], erow[7], erow[9]]
        if [norm(v) for v in actual] != [norm(v) for v in wanted]:
            errors.append(f'Element workbook row {n} differs from model')
        actual = list(supply.get(n, [None] * 6)[1:6])
        if [norm(v) for v in actual] != [norm(v) for v in wanted]:
            errors.append(f'Supply workbook row {n} differs from model')
        if e['model'] and norm(e['model']) not in norm(''.join(pdf_text['elements_pdf'])):
            errors.append(f'Model not found in element PDF: {n} {e["model"]}')
        if e['model'] and norm(e['model']) not in norm(''.join(pdf_text['equipment_pdf'])):
            errors.append(f'Model not found in supply PDF: {n} {e["model"]}')

    conn = [r for r in rows(path('mounting_workbook'), 'Соединения') if r and type(r[0]) is int]
    if len(conn) != len(model['connections']):
        errors.append('Connection workbook count differs')
    for a, b in zip(conn, model['connections']):
        wanted = [b[k] for k in ('num','net','a','b','section','color','zone')] + [b.get('material') or b['kind']]
        if [norm(v) for v in a[:8]] != [norm(v) for v in wanted]:
            errors.append(f'Connection workbook row {b["num"]} differs')
    def edge(net, a, b):
        return (str(net), *sorted([a, b]))
    a = Counter(edge(e['net'], e['a'], e['b']) for e in elec)
    b = Counter(edge(e['net'], e['a_raw'], e['b_raw']) for e in model['connections'])
    graph_equal = a == b
    if not graph_equal:
        errors.append({'electrical_to_mounting_difference': {'only_electrical':list((a-b).elements()),'only_mounting':list((b-a).elements())}})

    terms = rows(path('mounting_workbook'), 'Клеммы')[4:]
    terms = [r for r in terms if r and r[0]]
    if len(terms) != len(model['terminals']):
        errors.append('Terminal workbook count differs')
    for a, b in zip(terms, model['terminals']):
        wanted = [b.get(k, '') for k in ('zone','terminal','net','model','local','external','cable')]
        if [norm(v) for v in a[:7]] != [norm(v) for v in wanted]:
            errors.append(f'Terminal workbook row {b["zone"]}.{b["terminal"]} differs')

    for role, page, tokens in [
        ('mounting', 0, ['R5ST1284', '1200×800×400', '1150×750', 'КВ3', '270']),
        ('mounting', 1, ['R5ST0542', '500×400×200', 'HL1', 'HL2', 'SB0']),
        ('elements_pdf', 0, ['R5ST1284', '1200×800×400', 'R5ST0542', '500×400×200']),
        ('equipment_pdf', 0, ['R5ST1284', '1200×800×400', 'R5ST0542', '500×400×200']),
    ]:
        text = norm(pdf_text[role][page])
        for token in tokens:
            if norm(token) not in text:
                errors.append(f'Missing cabinet/layout evidence {role} page {page+1}: {token}')
    for role in ['mounting','elements_pdf','equipment_pdf']:
        if 'R5ST1084' in ''.join(pdf_text[role]) or '1000×800×400' in norm(''.join(pdf_text[role])):
            errors.append(f'Obsolete power enclosure in {role}')
    for n, article, quantity in [(1,'R5ST1284',2),(2,'R5ST0542',2)]:
        if spec[n][2] != article or spec[n][4] != quantity or supply[n][4] != quantity:
            errors.append(f'Cabinet procurement mismatch: {n}')
    placed = {x['tag']:x for x in layout['placements']}
    ducts = {x['tag']:x for x in layout['ducts']}
    rails = {x['tag']:x for x in layout['rails']}
    layout_matches = (placed['UZ1']['x'], placed['UZ1']['y'], rails['DIN2']['w']) == (550,373,430)
    duct_matches = (ducts['КВ3']['x'],ducts['КВ3']['y'],ducts['КВ3']['h']) == (480,220,410)
    if not layout_matches:
        errors.append('Wrong R31 panel geometry revision')
    if not duct_matches:
        errors.append('Missing R31 vertical connecting duct')

    # Read quantities back from the newly rendered supply PDF rows.
    cells_path = root / 'model/equipment_pdf_cells.json'
    if cells_path.is_file():
        with fitz.open(path('equipment_pdf')) as pdf:
            cells = json.loads(cells_path.read_text())
            expected_keys = {(str(n), col) for n in range(1,100) for col in range(4)} | {('—',col) for col in range(4)}
            actual_keys = {(str(c['row']),c['col']) for c in cells}
            if len(cells) != 400 or actual_keys != expected_keys:
                errors.append('Incomplete or duplicate supply PDF cell evidence')
            by_position = {str(e['spec']):e for e in model['elements']}
            for cell in cells:
                e = by_position.get(str(cell['row']))
                if e and cell['col'] == 2:
                    expected_quantity = str(e['purchase']).replace('.',',') + ' ' + e['unit']
                    if norm(cell['expected']) != norm(expected_quantity):
                        errors.append(f'Supply PDF quantity differs from model: {cell["row"]}')
                rect = fitz.Rect([v*72/25.4 for v in cell['bbox_mm']])
                text = norm(pdf[cell['page']-1].get_textbox(rect))
                if text != norm(cell['expected']):
                    errors.append(f'Supply PDF readback row {cell["row"]}, col {cell["col"]}')
            evidence['supply_pdf_cells_read_back'] = len(cells)
    else:
        errors.append('Missing supply PDF readback evidence')
    evidence.update(documents=len(docs),pdf_pages=sum(d.get('pages',0) for d in docs.values()),
                    element_rows=len(spec),connections=len(conn),terminals=len(terms),
                    electrical_connections_equal=graph_equal,
                    cabinet_dimensions=manifest['cabinets'],R31_geometry=layout_matches and duct_matches)
    return {'reference_id':manifest['id'],'passed_within_scope':not errors,
            'checked':evidence,'errors':errors,'known_limitations':manifest['limitations'],
            'scope':'Frozen-file integrity; 99 element and purchase rows; workbook connections/terminals; electrical graph equality; cabinet dimensions; R31 panel identification; supply PDF cell readback. Not a new electrical/thermal/safety approval.'}


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('root', type=Path)
    ap.add_argument('--report', type=Path)
    args = ap.parse_args()
    result = check(args.root)
    if args.report:
        args.report.parent.mkdir(parents=True,exist_ok=True)
        args.report.write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False,indent=2))
    raise SystemExit(0 if result['passed_within_scope'] else 2)
